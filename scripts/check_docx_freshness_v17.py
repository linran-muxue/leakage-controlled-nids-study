"""Check that the DOCX deliverables carry the current manuscript content.
The DOCX files are what gets submitted and what the advisor reads, but they are
generated artifacts: editing the Markdown and forgetting to rebuild leaves a
stale DOCX behind. This check reads the DOCX text and looks for the markers
that only the current revision contains.
"""
from __future__ import annotations
import sys
from pathlib import Path
from docx import Document
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "重构版论文_v4_20260915"
EXPECTED = {
    "English_SCI_Manuscript_v4.docx": ["v1.10.0", "3,469", "aggregation-rule differences"],
    "中文SCI论文_v4_重构版.docx": ["v1.10.0", "3469", "聚合策略差异"],
    "论文自查表.docx": ["E9", "补充材料_S01_S26"],
    "遗漏问题审查报告.docx": ["第十一轮"],
    "论文结构诊断与重构方案.docx": ["状态说明"],
    "研究缺口审计与优先级清单.docx": ["状态说明"],
    "P0_P1执行手册.docx": ["状态说明"],
}
def docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
def main() -> int:
    problems: list[str] = []
    for name, markers in EXPECTED.items():
        path = BASE / name
        if not path.exists():
            problems.append(f"{name} is missing")
            print(f"  {name}: MISSING")
            continue
        text = docx_text(path)
        missing = [m for m in markers if m not in text]
        print(f"  {name}: {len(text)} chars, missing markers {missing or 'none'}")
        if missing:
            problems.append(f"{name} is out of date (missing {missing}); rebuild it from the Markdown")
    print()
    if problems:
        for problem in problems:
            print(f"ISSUE {problem}")
        print("DOCX_STALE")
        return 1
    print("DOCX_OK")
    return 0
if __name__ == "__main__":
    raise SystemExit(main())
