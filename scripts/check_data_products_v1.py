"""Verify the data-products inventory against the local data.

The document is a claim about files that are deliberately not shipped, so the
strongest check is to regenerate it: the builder walks the data directories,
counts rows, hashes every file up to 64 MB and asserts the published population
sizes.  A byte-identical regeneration means every number in the document still
describes the data on disk.

Rebuild with ``scripts/build_data_products_doc_v1.py`` and
``scripts/build_restructured_manuscript_v4.py --only dataproducts`` on a mismatch.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "重构版论文_v4_20260915"
DOC = BASE / "数据产物清单.md"
WORD = BASE / "数据产物清单.docx"
ANCHORS = ("数据产物清单", "论文用到的总体", "去重与清洗的阶段计数", "重建命令")


def main() -> int:
    from docx import Document

    problems: list[str] = []
    for path in (DOC, WORD):
        if not path.exists():
            problems.append(f"missing artefact: {path.name}")
    if problems:
        print("DATA_PRODUCTS_FAILED: " + "; ".join(problems))
        return 1

    directories = sorted(p.name for p in ROOT.iterdir()
                         if p.is_dir() and (p.name.startswith("data_processed")
                                            or p.name.startswith("data_external")))
    text = DOC.read_text(encoding="utf-8")
    for name in directories:
        if f"`{name}`" not in text:
            problems.append(f"data directory missing from the inventory: {name}")
    if not problems:
        with tempfile.TemporaryDirectory() as tmp:
            proc = subprocess.run([sys.executable,
                                   str(ROOT / "scripts" / "build_data_products_doc_v1.py"),
                                   "--outdir", tmp], cwd=ROOT, capture_output=True, text=True,
                                  errors="replace")
            if proc.returncode != 0:
                problems.append("the builder failed: "
                                + " ".join((proc.stdout + proc.stderr).split())[-240:])
            else:
                fresh = (Path(tmp) / "数据产物清单.md").read_text(encoding="utf-8")
                if fresh != text:
                    problems.append("数据产物清单.md is stale: rebuild it with "
                                    "scripts/build_data_products_doc_v1.py")

    document = Document(str(WORD))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                body += "\n" + cell.text
    for anchor in ANCHORS:
        if anchor not in body:
            problems.append(f"数据产物清单.docx lost the anchor {anchor!r}")

    if problems:
        print("DATA_PRODUCTS_FAILED: " + "; ".join(problems[:6]))
        return 1
    print(f"DATA_PRODUCTS_OK directories={len(directories)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
