from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "agent_paper_key_points.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.add_run(item)


def add_paper(doc, number, title, question, mechanism, keywords, security):
    doc.add_heading(f"{number}. {title}", level=1)
    p = doc.add_paragraph()
    p.add_run("研究问题：").bold = True
    p.add_run(question)
    p = doc.add_paragraph()
    p.add_run("核心机制：").bold = True
    p.add_run(mechanism)
    p = doc.add_paragraph()
    p.add_run("关键词：").bold = True
    p.add_run(keywords)
    p = doc.add_paragraph()
    p.add_run("安全方向启发：").bold = True
    p.add_run(security)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 15, "1F4E79"), ("Heading 2", 12, "2F75B5")]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("LLM Agent 论文重点知识点")
    r.bold = True
    r.font.size = Pt(21)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("ReAct、Toolformer、Reflexion 等核心论文精简笔记")
    sr.font.size = Pt(11)
    sr.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_heading("一、Agent 必备基本概念", level=1)
    add_bullets(doc, [
        "Agent：能够在环境中观察、决策、行动并根据反馈继续完成任务的系统。",
        "Environment：Agent 工作的环境，如网页、数据库、代码仓库或安全日志系统。",
        "Observation：Agent 当前获得的信息；Action：Agent 执行的动作。",
        "Tool：Agent 可调用的外部函数、API、搜索器或数据库。",
        "State：环境当前状态；Policy：根据历史信息选择下一步动作的策略。",
        "Memory：保存历史信息、经验和任务进度；Trajectory：完整行动轨迹。",
        "Evaluator：判断任务是否完成的评测程序。",
        "核心闭环：任务 → 观察 → 推理/规划 → 工具调用 → 反馈 → 继续或结束。",
    ])

    doc.add_heading("二、八篇论文核心知识点", level=1)
    add_paper(doc, 1, "ReAct", "如何让语言模型边推理边行动？", "Thought → Action → Observation 循环；根据工具返回结果决定下一步。", "推理、行动、观察、轨迹、工具调用", "安全 Agent 可先分析告警，再查询 IP、流量和日志，并根据结果更新判断。")
    add_paper(doc, 2, "Toolformer", "模型何时调用工具、调用什么工具以及传递什么参数？", "学习在文本生成过程中插入 API 调用，并融合工具返回结果。", "API、工具选择、参数生成、结果融合", "研究安全 Agent 如何选择最合适的查询工具，避免无效或过度调用。")
    add_paper(doc, 3, "Reflexion", "Agent 如何利用失败经验改进下一次尝试？", "任务失败后生成语言反思，保存为记忆，在下一次任务中读取。", "反馈、反思、情景记忆、自我改进", "可研究告警分析失败后的反思是否提高准确率，以及额外 Token 成本。")
    add_paper(doc, 4, "Tree of Thoughts", "如何探索、评估和比较多条推理路径？", "生成多个候选思路，进行评估、剪枝和回溯，而不是只走一条推理链。", "搜索、候选路径、评估、剪枝、回溯", "面对复杂攻击链，可比较多个调查方案后再确定风险结论。")
    add_paper(doc, 5, "Generative Agents", "如何让 Agent 具备长期记忆、反思和计划？", "记录经历，按相关性/重要性/新近性检索，形成反思并用于计划。", "记忆流、检索、反思、长期计划", "可保存 IP、主机和用户的历史行为，但必须防止错误或过时记忆污染判断。")
    add_paper(doc, 6, "AgentBench", "如何在多种环境中系统评估 Agent？", "用数据库、知识图谱、操作系统、网页等环境测试任务完成和交互能力。", "Benchmark、多环境、任务成功、自动评测", "安全 Agent 必须预先定义环境、成功标准和评测器，不能只展示案例。")
    add_paper(doc, 7, "τ-bench", "Agent 能否在真实业务规则下正确使用工具？", "模拟业务对话，检查工具调用、规则遵循、状态保持和权限边界。", "工具调用、对话状态、业务规则、权限", "安全 Agent 只能查询授权数据，危险操作应要求人工确认。")
    add_paper(doc, 8, "AgentDojo", "Agent 如何抵抗提示注入并保持正常任务能力？", "在动态工具环境中同时评估正常任务效用和攻击防御能力。", "提示注入、间接注入、效用、安全、越权", "日志、网页和威胁情报可能含恶意指令，应隔离数据与命令并限制权限。")

    doc.add_heading("三、八篇论文的整体逻辑", level=1)
    p = doc.add_paragraph()
    p.add_run("ReAct：").bold = True
    p.add_run("让 Agent 能行动；")
    p.add_run("Toolformer：").bold = True
    p.add_run("让模型学会使用工具；")
    p.add_run("Reflexion：").bold = True
    p.add_run("让 Agent 从失败中改进；")
    p.add_run("Tree of Thoughts：").bold = True
    p.add_run("让 Agent 搜索多条方案；")
    p.add_run("Generative Agents：").bold = True
    p.add_run("让 Agent 拥有长期记忆；")
    p.add_run("AgentBench/τ-bench：").bold = True
    p.add_run("定义如何评测；")
    p.add_run("AgentDojo：").bold = True
    p.add_run("研究安全风险。")

    doc.add_heading("四、实验中必须掌握的指标", level=1)
    add_bullets(doc, [
        "任务成功率 / 准确率：最终任务是否完成、判断是否正确。",
        "工具调用次数和平均步骤数：Agent 是否高效。",
        "Token 消耗和响应延迟：Agent 的资源成本。",
        "工具错误率和重试率：行动是否可靠。",
        "失败类型：规划错误、参数错误、证据理解错误、权限错误、恢复失败。",
        "安全实验还要报告攻击成功率，以及正常任务成功率（效用—安全权衡）。",
    ])

    doc.add_heading("五、与网络安全告警分析的连接", level=1)
    add_bullets(doc, [
        "输入：入侵检测模型产生的告警。",
        "工具：IP 信誉查询、历史告警查询、流量统计、主机日志查询。",
        "输出：攻击/误报判断、攻击类型、风险等级、证据和处置建议。",
        "可研究变量：工具路由策略、反思频率、记忆方式或失败恢复策略。",
        "基本研究问题：在固定模型、工具和 Token 预算下，不同工具调用策略是否提高准确率并降低成本？",
    ])

    doc.add_heading("六、最先记住的三句话", level=1)
    add_bullets(doc, [
        "Agent 不只是生成文本，而是在环境中连续观察、决策和行动。",
        "科研重点不是做一个 Demo，而是提出一个可控制变量并用指标验证。",
        "安全 Agent 必须同时考虑任务成功率、调用成本和越权/提示注入风险。",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
