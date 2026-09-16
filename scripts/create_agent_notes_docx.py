from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "agent_paper_notes_copyable.md"
OUTPUT = ROOT / "docs" / "agent_paper_notes_copyable.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text.strip())
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def parse_table(lines, i, doc):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    if not rows:
        return i
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            value = row[ci] if ci < len(row) else ""
            set_cell_text(table.cell(ri, ci), value, bold=(ri == 0))
            if ri == 0:
                set_cell_shading(table.cell(ri, ci), "D9EAF7")
    doc.add_paragraph()
    return i


def build_doc():
    md = SOURCE.read_text(encoding="utf-8")
    lines = md.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2F75B5"), ("Heading 3", 11, "4F81BD")]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("LLM Agent 论文阅读笔记")
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "Microsoft YaHei"
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("ReAct、Toolformer、Reflexion 等核心论文")
    sr.font.size = Pt(12)
    sr.font.color.rgb = RGBColor(89, 89, 89)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dr = date.add_run("用于 Agent 科研入门与网络安全告警分析方向准备")
    dr.font.size = Pt(10)
    dr.font.color.rgb = RGBColor(127, 127, 127)
    doc.add_page_break()

    toc_title = doc.add_paragraph("目录", style="Heading 1")
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.right_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                for idx, code_line in enumerate(code_lines):
                    r = p.add_run(code_line)
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(64, 64, 64)
                    if idx < len(code_lines) - 1:
                        r.add_break()
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            i = parse_table(lines, i, doc)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_paragraph(m.group(2).strip(), style=f"Heading {level}")
            i += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(bullet.group(1).strip())
            i += 1
            continue
        numbered = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.add_run(numbered.group(2).strip())
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(line.strip())
        i += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LLM Agent 论文阅读笔记  |  ")
    add_field(footer, "PAGE")
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
