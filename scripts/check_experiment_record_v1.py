"""Verify the experiment record: it regenerates byte for byte, and its copies match.

The record is a deliverable that claims three things - these are the scripts, this
is what they produced, and this is the run evidence - so all three are checked:
the document and its eighteen panels are re-rendered into a scratch directory and
compared with the committed files, the copies in ``实验代码/`` are compared with
the repository originals, and the ``实验日志/`` copies with the logs they came
from where those still exist.  The Word copy must carry the document's anchors.

Rebuild with ``scripts/build_experiment_record_v1.py`` and then
``scripts/build_restructured_manuscript_v4.py --only experiments`` on a mismatch.
"""
from __future__ import annotations

import hashlib
import subprocess
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "重构版论文_v4_20260915"
PY = sys.executable
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document  # noqa: E402

DOC = BASE / "实验代码与运行记录.md"
WORD = BASE / "实验代码与运行记录.docx"
FIGDIR = BASE / "figures_experiments"
CODEDIR = BASE / "实验代码"
LOGDIR = BASE / "实验日志"
ANCHORS = ("实验代码与运行记录", "非屏幕截图", "exp05", "exp18", "附录 C")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    problems: list[str] = []
    for path in (DOC, WORD):
        if not path.exists():
            problems.append(f"missing artefact: {path.name}")
    panels = sorted(FIGDIR.glob("exp*.png")) if FIGDIR.exists() else []
    if len(panels) < 18:
        problems.append(f"only {len(panels)} run-record panels, expected 18")
    if problems:
        print("EXPERIMENT_RECORD_FAILED: " + "; ".join(problems))
        return 1

    with tempfile.TemporaryDirectory() as tmp:
        scratch = Path(tmp)
        proc = subprocess.run([PY, str(ROOT / "scripts" / "build_experiment_record_v1.py"),
                               "--outdir", tmp, "--figdir", str(scratch / "figures")],
                              cwd=ROOT, capture_output=True, text=True, errors="replace")
        if proc.returncode != 0:
            problems.append("the builder failed: " + " ".join((proc.stdout + proc.stderr).split())[-240:])
        else:
            if (scratch / "实验代码与运行记录.md").read_text("utf-8") != DOC.read_text("utf-8"):
                problems.append("实验代码与运行记录.md is stale: rebuild it with "
                                "scripts/build_experiment_record_v1.py")
            for panel in panels:
                fresh = scratch / "figures" / panel.name
                if not fresh.exists() or fresh.read_bytes() != panel.read_bytes():
                    problems.append(f"{panel.name} is stale: rebuild it with "
                                    "scripts/build_experiment_record_v1.py")

    # the copies must be byte-identical to the files they claim to copy
    text = DOC.read_text(encoding="utf-8")
    cited = sorted({name for name in _code_paths(text)})
    for relative in cited:
        source = ROOT / relative
        copy = CODEDIR / Path(relative).name
        if not source.exists():
            problems.append(f"cited script does not exist: {relative}")
        elif not copy.exists():
            problems.append(f"missing copy: {copy.name}")
        elif digest(source) != digest(copy):
            problems.append(f"copy differs from the repository file: {copy.name}")
    extras = sorted(p.name for p in CODEDIR.glob("*.py")
                    if f"/{p.name}" not in text and p.name not in text)
    if extras:
        problems.append(f"unreferenced scripts in 实验代码/: {', '.join(extras[:5])}")

    for copy in sorted(LOGDIR.glob("*.log")) if LOGDIR.exists() else []:
        source = ROOT / "logs" / copy.name
        if source.exists() and digest(source) != digest(copy):
            problems.append(f"log copy differs from logs/{copy.name}")

    document = Document(str(WORD))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                body += "\n" + cell.text
    for anchor in ANCHORS:
        if anchor not in body:
            problems.append(f"实验代码与运行记录.docx lost the anchor {anchor!r}")

    if problems:
        print("EXPERIMENT_RECORD_FAILED: " + "; ".join(problems))
        return 1
    print(f"EXPERIMENT_RECORD_OK experiments=18 panels={len(panels)} scripts={len(cited)} "
          f"logs={len(list(LOGDIR.glob('*.log'))) if LOGDIR.exists() else 0}")
    return 0


def _code_paths(text: str) -> list[str]:
    """Repository-relative script paths the document cites in its tables."""
    import re
    return re.findall(r"\| `((?:scripts|src)/[^`]+\.py)` \|", text)


if __name__ == "__main__":
    sys.exit(main())
