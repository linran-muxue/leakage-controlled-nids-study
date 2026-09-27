"""Geometric fit check for the briefing deck.

PowerPoint's own renderer is the reference for text fit, but it cannot be driven
from this session (the Office COM preflight refuses a non-active desktop and no
alternate renderer is installed).  This check therefore measures the deck
geometrically: every text frame is laid out with an estimated glyph advance
(full-width characters one em, the rest about 0.55 em), wrapped at the shape's
inner width, and compared with the height it has.  It also rejects shapes that
leave the slide and text frames that overlap each other.

It is a risk filter, not a substitute for opening the deck: anything it passes
still deserves one look in PowerPoint.
"""
from __future__ import annotations

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "重构版论文_v4_20260915" / "汇报用_论文介绍.pptx"
PT = 12700  # EMU per point
problems: list[str] = []
checked = 0


def advance(character: str, size: float) -> float:
    """Estimated glyph advance in points."""
    code = ord(character)
    if code > 0x2E80:  # CJK, fullwidth forms and punctuation
        return size
    if character.isupper() or character in "mw":
        return size * 0.62
    return size * 0.52


def frame_size(shape) -> tuple[float, float]:
    frame = shape.text_frame
    left = frame.margin_left / PT if frame.margin_left is not None else 7.2
    right = frame.margin_right / PT if frame.margin_right is not None else 7.2
    top = frame.margin_top / PT if frame.margin_top is not None else 3.6
    bottom = frame.margin_bottom / PT if frame.margin_bottom is not None else 3.6
    return (shape.width / PT - left - right, shape.height / PT - top - bottom)


def needed_height(shape) -> float:
    inner_w, _ = frame_size(shape)
    total = 0.0
    for paragraph in shape.text_frame.paragraphs:
        runs = paragraph.runs
        text = "".join(run.text for run in runs)
        if not text:
            total += 6.0
            continue
        size = max(float(run.font.size.pt) for run in runs if run.font.size) if any(
            run.font.size for run in runs) else 18.0
        width = sum(advance(character, size) for character in text)
        lines = max(1, math.ceil(width / inner_w)) if inner_w > 0 else 1
        spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.2
        total += lines * size * 1.2 * spacing
    return total


def main() -> int:
    global checked
    deck = Presentation(str(DECK))
    slide_w, slide_h = deck.slide_width / PT, deck.slide_height / PT
    for index, slide in enumerate(deck.slides, 1):
        frames = []
        for shape in slide.shapes:
            if shape.left is None:
                continue
            x, y = shape.left / PT, shape.top / PT
            w, h = (shape.width or 0) / PT, (shape.height or 0) / PT
            if x < -1 or y < -1 or x + w > slide_w + 1 or y + h > slide_h + 1:
                problems.append(f"slide {index}: a shape leaves the slide "
                                f"({x:.0f},{y:.0f} {w:.0f}x{h:.0f})")
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            checked += 1
            _, inner_h = frame_size(shape)
            needed = needed_height(shape)
            if needed > inner_h + 1:
                problems.append(f"slide {index}: '{shape.text_frame.text[:28]}…' needs "
                                f"{needed:.0f}pt in a {inner_h:.0f}pt box")
            frames.append((x, y, w, h, shape.text_frame.text[:24]))
        for i in range(len(frames)):
            for j in range(i + 1, len(frames)):
                ax, ay, aw, ah, atext = frames[i]
                bx, by, bw, bh, btext = frames[j]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 6 and oy > 6:
                    problems.append(f"slide {index}: '{atext}' overlaps '{btext}' "
                                    f"({ox:.0f}x{oy:.0f}pt)")
    print(f"text frames measured: {checked}; slides: {len(deck.slides._sldIdLst)}")
    print(f"problems: {len(problems)}")
    for problem in problems:
        print(f"  ISSUE {problem}")
    if problems:
        print("DECK_FIT_FAILED")
        return 1
    print("DECK_FIT_OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
