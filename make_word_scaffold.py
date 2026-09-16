from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path.cwd()
papers=[
 ('2303.11366v4.pdf','Reflexion：具有言语强化学习的语言智能体','paper1_en.txt','大型语言模型（LLM）正越来越多地被用于与外部环境（如游戏、编译器和 API）交互，充当目标驱动的智能体。然而，这些语言智能体很难像传统强化学习方法那样通过试错快速而高效地学习，因为传统方法需要大量训练样本以及昂贵的模型微调。我们提出 Reflexion，一种不通过更新模型权重、而是通过语言反馈来强化语言智能体的新框架。具体而言，Reflexion 智能体会对任务反馈信号进行语言化反思，并将反思文本保存在情景记忆缓冲区中，以促使后续试验作出更好的决策。该方法能够灵活整合不同类型（标量值或自由形式语言）和不同来源（外部或内部模拟）的反馈信号，并在顺序决策、编程和语言推理等任务上显著优于基线。') ,
 ('2210.03629v3.pdf','ReAct：在语言模型中协同推理与行动','paper2_en.txt','尽管大型语言模型在语言理解和交互式决策任务上表现出色，但其推理能力（例如思维链提示）与行动能力（例如行动计划生成）主要被作为彼此独立的主题研究。本文探索让语言模型以交错方式同时生成推理轨迹和任务特定行动，从而实现二者更紧密的协同：推理轨迹帮助模型归纳、跟踪和更新行动计划，并处理异常情况；行动则使模型能够与外部信息源（如知识库或环境）交互并收集额外信息。我们将这种方法命名为 ReAct，并在多种语言和决策任务上验证其有效性。') ,
 ('2302.04761v1.pdf','Toolformer：语言模型能够自我学习使用工具','paper3_en.txt','语言模型能够仅凭少量示例或文本指令解决新任务，尤其是在模型规模较大时表现突出。然而，它们有时仍难以完成算术或事实查找等基础功能，而更简单、更小的模型反而擅长这些任务。本文展示了语言模型如何通过简单 API 自我学习使用外部工具，从而兼具两方面优势。我们提出 Toolformer：一种能够学习决定何时调用 API、调用哪些 API、传递什么参数，以及如何将结果整合到后续词元预测中的模型。该过程以自监督方式完成，每个 API 只需少量示例。')
]
doc=Document(); doc.styles['Normal'].font.name='等线'; doc.styles['Normal'].font.size=Pt(10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('三篇人工智能论文中文翻译文档'); r.bold=True; r.font.size=Pt(18)
doc.add_paragraph('说明：已导入三篇 PDF 的完整可提取文本，并提供中文标题与摘要译文。由于当前翻译服务不可用，正文尚未完成自动翻译；英文原文已按论文分隔保留，便于后续继续翻译和校对。')
for fname,title,enfile,abstract in papers:
 doc.add_page_break(); h=doc.add_heading(title,level=1); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.add_paragraph('原文件：'+fname)
 doc.add_heading('摘要（中文译文）',level=2); doc.add_paragraph(abstract)
 doc.add_heading('英文原文（完整提取）',level=2)
 txt=(ROOT/enfile).read_text(encoding='utf-8')
 for block in txt.split('\n\n'):
  if block.strip(): doc.add_paragraph(block.strip())
out=ROOT/'三篇论文中文翻译.docx'; doc.save(out); print(out)
